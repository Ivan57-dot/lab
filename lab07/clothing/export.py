from docx import Document
from openpyxl import Workbook


class Exporter:
    
    def __init__(self):
        self.n = 0
    
    def docx(self, file, name, size, fabric, cost):
        d = Document()
        d.add_heading('Расчёт', 0)
        d.add_paragraph(f'Изделие: {name}')
        d.add_paragraph(f'Размер: {size}')
        d.add_paragraph(f'Ткань: {fabric} м')
        d.add_paragraph(f'Стоимость: {cost} руб')
        d.save(file)
        self.n += 1
    
    def xlsx(self, file, name, size, fabric, cost):
        w = Workbook()
        s = w.active
        s['A1'], s['B1'] = 'Параметр', 'Значение'
        s['A2'], s['B2'] = 'Изделие', name
        s['A3'], s['B3'] = 'Размер', size
        s['A4'], s['B4'] = 'Ткань (м)', fabric
        s['A5'], s['B5'] = 'Стоимость (руб)', cost
        w.save(file)
        self.n += 1
    
    def __len__(self):
        return self.n
    
    def __str__(self):
        return f"Сохранено: {self.n}"