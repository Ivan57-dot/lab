from docx import Document
from openpyxl import Workbook
def save_to_docx(filename, chto, razmer, fabric, cost):
    doc = Document()
    doc.add_heading('Результат расчёта', 0)
    doc.add_paragraph('Изделие: ' + chto)
    doc.add_paragraph('Размер: ' + str(razmer))
    doc.add_paragraph('Ткань: ' + str(fabric) + ' м')
    doc.add_paragraph('Стоимость: ' + str(cost) + ' руб')
    doc.save(filename)
def save_to_xlsx(filename, chto, razmer, fabric, cost):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = 'Параметр'
    ws['B1'] = 'Значение'
    ws['A2'] = 'Изделие'
    ws['B2'] = chto
    ws['A3'] = 'Размер'
    ws['B3'] = razmer
    ws['A4'] = 'Ткань (м)'
    ws['B4'] = fabric
    ws['A5'] = 'Стоимость (руб)'
    ws['B5'] = cost
    wb.save(filename)